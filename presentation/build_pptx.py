"""Generates a Google Slides-importable .pptx of the Module 3 ADK lab talk.

Google Slides has no native HTML or Markdown import, but it imports .pptx with high
fidelity. This builds the same 15 slides as the reveal.js deck using native PowerPoint
shapes and tables -- not images -- so every element stays editable after import.

The talk track goes into the PowerPoint notes field, which Slides maps onto its own
speaker notes.

Run with:  uvx --with python-pptx python build_pptx.py
"""

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = pathlib.Path(__file__).parent / "cymbal-agent-module3.pptx"

# GitHub-dark palette, matching the reveal deck.
BG      = RGBColor(0x0D, 0x11, 0x17)
CARD    = RGBColor(0x16, 0x1B, 0x22)
LINE    = RGBColor(0x30, 0x36, 0x3D)
FG      = RGBColor(0xE6, 0xED, 0xF3)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
MUTED   = RGBColor(0x8B, 0x94, 0x9E)
BLUE    = RGBColor(0x58, 0xA6, 0xFF)
GREEN   = RGBColor(0x3F, 0xB9, 0x50)
AMBER   = RGBColor(0xD2, 0x99, 0x22)
RED     = RGBColor(0xF8, 0x51, 0x49)
PURPLE  = RGBColor(0xBC, 0x8C, 0xFF)

SANS = "Arial"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def slide(prs, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if notes:
        s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def textbox(s, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=18, color=FG, bold=False, font=SANS,
         align=None, space_after=6, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return p


def rich(tf, segments, size=18, align=None, space_after=6, first=False):
    """Adds a paragraph built from (text, color, bold, font) tuples."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    for seg in segments:
        text, color = seg[0], seg[1]
        bold = seg[2] if len(seg) > 2 else False
        font = seg[3] if len(seg) > 3 else SANS
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return p


def kicker(s, text):
    tf = textbox(s, Inches(0.7), Inches(0.45), Inches(12), Inches(0.4))
    para(tf, text.upper(), size=12, color=MUTED, bold=True, first=True)


def title(s, text, y=0.85, size=34, color=WHITE, align=PP_ALIGN.LEFT):
    tf = textbox(s, Inches(0.7), Inches(y), Inches(12), Inches(0.9), align=align)
    para(tf, text, size=size, color=color, bold=True, align=align, first=True)


def card(s, x, y, w, h, accent, heading, body_segments, metric=None):
    """A rounded panel with a coloured left edge."""
    box = s.shapes.add_shape(5, x, y, w, h)  # ROUNDED_RECTANGLE
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = LINE
    box.line.width = Pt(0.75)
    box.shadow.inherit = False

    bar = s.shapes.add_shape(1, x, y, Emu(38100), h)  # RECTANGLE, ~3pt wide
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False

    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_right = Inches(0.18), Inches(0.12)
    tf.margin_top, tf.margin_bottom = Inches(0.12), Inches(0.1)
    para(tf, heading, size=14, color=WHITE, bold=True, space_after=4, first=True)
    if body_segments:
        rich(tf, body_segments, size=11, space_after=3)
    if metric:
        para(tf, metric, size=11, color=GREEN, font=MONO, space_after=0)
    return box


def callout(s, x, y, w, h, segments, accent=BLUE, size=13):
    box = s.shapes.add_shape(5, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = accent
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left, tf.margin_right = Inches(0.2), Inches(0.2)
    rich(tf, segments, size=size, space_after=0, first=False)
    # remove the empty seed paragraph
    tf.paragraphs[0]._p.getparent().remove(tf.paragraphs[0]._p)
    return box


def table(s, x, y, w, rows, col_widths=None, header=True, font_size=12):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = s.shapes.add_table(n_rows, n_cols, x, y, w, Inches(0.4 * n_rows))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = cw
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(0.36)
        for ci, cell_spec in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG if ri else CARD
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            if isinstance(cell_spec, tuple):
                text, color = cell_spec[0], cell_spec[1]
                fnt = cell_spec[2] if len(cell_spec) > 2 else SANS
                algn = cell_spec[3] if len(cell_spec) > 3 else PP_ALIGN.LEFT
            else:
                text, color, fnt, algn = cell_spec, FG, SANS, PP_ALIGN.LEFT

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.alignment = algn
            for r in p.runs:
                r.font.size = Pt(font_size)
                r.font.name = fnt
                r.font.bold = bool(header and ri == 0)
                r.font.color.rgb = MUTED if (header and ri == 0) else color
    return tbl


def code_block(s, x, y, w, h, lines):
    box = s.shapes.add_shape(5, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = LINE
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.3)
    for i, line in enumerate(lines):
        para(tf, line, size=18, color=BLUE, font=MONO, space_after=2, first=(i == 0))
    return box


def chip(s, x, y, w, h, top, bottom, accent=LINE):
    box = s.shapes.add_shape(5, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = CARD
    box.line.color.rgb = accent
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, top, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
         space_after=1, first=True)
    para(tf, bottom, size=9, color=MUTED, font=MONO, align=PP_ALIGN.CENTER,
         space_after=0)
    return box


def arrow(s, x, y):
    tf = textbox(s, x, y, Inches(0.4), Inches(0.4), align=PP_ALIGN.CENTER)
    para(tf, "\u2192", size=20, color=MUTED, align=PP_ALIGN.CENTER, first=True)


# ─────────────────────────────────────────────────────────────────────────────
prs = new_deck()

# 1 · TITLE
s = slide(prs, """
8 minutes. Don't introduce yourself for more than one line - get to the operator's
question fast. Thesis to land early: THE MODEL IS A ROUTER; the hard parts were
infrastructure.
""")
tf = textbox(s, Inches(1), Inches(2.4), Inches(11.3), Inches(3), align=PP_ALIGN.CENTER)
para(tf, "DATA ANALYTICS ADVANCED ELEVATE  ·  MODULE 3", size=13, color=MUTED,
     bold=True, align=PP_ALIGN.CENTER, space_after=18, first=True)
para(tf, "Cymbal Operations Agent", size=52, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, space_after=14)
para(tf, "One question. Three data systems. No humans in between.",
     size=20, color=BLUE, align=PP_ALIGN.CENTER, space_after=30)
para(tf, "Google ADK  ·  Gemini 3.6 Flash  ·  BigQuery  ·  Bigtable  ·  Cloud Run",
     size=12, color=MUTED, font=MONO, align=PP_ALIGN.CENTER)

# 2 · THE PROBLEM
s = slide(prs, """
45 seconds. Say the question out loud as an operator would. Then: "that's three systems,
three teams, three days." Pause. Don't rush into architecture.
""")
kicker(s, "The question nobody can answer")
tf = textbox(s, Inches(0.7), Inches(1.0), Inches(12), Inches(1.9))
rich(tf, [
    ('"My Ginza store - is it about to ', FG),
    ("stock out", BLUE, True),
    (", is the ", FG),
    ("card reader at lane 3", BLUE, True),
    (" broken, and is ", FG),
    ("cashier 1190", BLUE, True),
    (' abusing manager overrides?"', FG),
], size=27, space_after=0)

cw, gap = Inches(3.85), Inches(0.28)
for i, (accent, head, body) in enumerate([
    (BLUE,   "Data warehouse",  "Inventory, sales, 7-day baselines"),
    (PURPLE, "A pile of PDFs",  "Vendor service manuals nobody reads"),
    (GREEN,  "Real-time store", "Live per-cashier counters"),
]):
    card(s, Inches(0.7) + i * (cw + gap), Inches(3.3), cw, Inches(1.15),
         accent, head, [(body, MUTED)])

tf = textbox(s, Inches(0.7), Inches(4.9), Inches(12), Inches(1.2))
para(tf, "Three systems. Three teams. Usually three days.",
     size=21, color=WHITE, bold=True, space_after=6, first=True)
para(tf, "Nobody owns the question, because the question crosses all three.",
     size=15, color=MUTED)

# 3 · ARCHITECTURE
s = slide(prs, """
60 seconds. Three beats:
(1) One model, four tools - it routes, it doesn't store.
(2) Latencies differ by an ORDER OF MAGNITUDE - 2s to 18s. That's a design constraint;
    it dictates what fits in one turn.
(3) Every tool is an ordinary service with ordinary IAM. That's the point.

If asked about security: the Bigtable tool is a PRIVATE Cloud Run service. The agent
impersonates a service account, mints an OIDC ID token scoped to the Cloud Run audience,
and calls with roles/run.invoker. No public endpoint, no API keys.
""")
kicker(s, "The whole system")
chip(s, Inches(4.9), Inches(1.0), Inches(3.5), Inches(0.75),
     "Coordinator Agent", "gemini-3.6-flash", accent=BLUE)

cw, gap = Inches(2.95), Inches(0.22)
for i, (accent, head, body, metric) in enumerate([
    (BLUE,   "Analytics",         "NL2SQL Data Agent\nBigQuery gold tables", "~18s"),
    (PURPLE, "Runbook RAG",       "VECTOR_SEARCH\n129 manual chunks",        "~2s"),
    (GREEN,  "Real-time",         "MCP on Cloud Run\nBigtable",              "~3s warm"),
    (AMBER,  "Entity resolution", "VECTOR_SEARCH\n16-store directory",       "~2s"),
]):
    card(s, Inches(0.7) + i * (cw + gap), Inches(2.15), cw, Inches(1.65),
         accent, head, [(body, MUTED)], metric=metric)

callout(s, Inches(0.7), Inches(4.2), Inches(11.95), Inches(1.0), [
    ("The model ", FG), ("writes no ungoverned SQL", WHITE, True),
    (" and ", FG), ("holds no data", WHITE, True),
    (". It picks a tool and passes arguments.\nEvery tool is an ordinary service with "
     "ordinary IAM.", FG),
], accent=BLUE, size=15)

# 4 · DEMO 1
s = slide(prs, """
TALK OVER THE 9 SECONDS. Don't watch the spinner.

Say: "This isn't the model remembering a manual. It's embedding the question, running a
cosine search over 129 chunks of a real Toshiba service PDF, and stitching neighbouring
chunks so the procedure isn't cut in half."

THEN POINT AT TWO THINGS:
1. "Do NOT re-swipe immediately" - a vendor safety instruction. If the model invented
   this it might say the opposite, and that's a duplicate charge on a customer's card.
2. The storage.cloud.google.com link - every answer carries its source PDF. An auditor
   can click through. Ungrounded output is not an answer, it's a suggestion.
""")
tf = textbox(s, Inches(0.7), Inches(1.3), Inches(12), Inches(1), align=PP_ALIGN.CENTER)
para(tf, "DEMO 1   ·   LIVE   ·   9.3s", size=13, color=GREEN, bold=True,
     align=PP_ALIGN.CENTER, space_after=14, first=True)
para(tf, "It works \u2014 and it cites", size=40, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER)
code_block(s, Inches(2.4), Inches(3.3), Inches(8.5), Inches(1.5),
           ["How do I fix error code ERR-PAY-4001",
            "on the POS EMV terminal reader?"])
tf = textbox(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.6), align=PP_ALIGN.CENTER)
para(tf, "Switch to  adk web", size=15, color=MUTED, font=MONO,
     align=PP_ALIGN.CENTER, first=True)

# 5 · DEMO 2 SETUP
s = slide(prs, """
THIS IS YOUR BEST 40 SECONDS. Set it up BEFORE you type:
"Now watch what happens when the question is unanswerable."
Fastest demo you have - 5.3 seconds.
""")
tf = textbox(s, Inches(0.7), Inches(1.3), Inches(12), Inches(1), align=PP_ALIGN.CENTER)
para(tf, "DEMO 2   ·   LIVE   ·   5.3s", size=13, color=GREEN, bold=True,
     align=PP_ALIGN.CENTER, space_after=14, first=True)
rich(tf, [("Now watch it ", WHITE, True), ("refuse", BLUE, True)],
     size=40, align=PP_ALIGN.CENTER)
code_block(s, Inches(2.4), Inches(3.3), Inches(8.5), Inches(1.5),
           ["What is the total on-hand inventory",
            "at the Ginza store?"])

# 6 · THE REFUSAL
s = slide(prs, """
Deliver the punchline slowly. Then the generalisation - this is the line to leave in
their heads.

ALSO SAY: the agent did NOT call the analytics tool at all here. The refusal is
structural, not cosmetic. There's a live integration test asserting exactly that.
""")
kicker(s, "Two stores. Identical names.")
table(s, Inches(0.7), Inches(1.05), Inches(11.95), [
    ["store_id", "store_name", "city", ("cosine", MUTED, SANS, PP_ALIGN.RIGHT)],
    [("STORE_001", BLUE, MONO), "Cymbal Tokyo Ginza District Flagship", "Tokyo",
     ("0.6955", FG, MONO, PP_ALIGN.RIGHT)],
    [("STORE_013", BLUE, MONO), "Cymbal Tokyo Ginza District Flagship", "Tokyo",
     ("0.6915", FG, MONO, PP_ALIGN.RIGHT)],
], col_widths=[Inches(2.0), Inches(6.35), Inches(2.0), Inches(1.6)], font_size=14)

callout(s, Inches(0.7), Inches(2.65), Inches(11.95), Inches(0.85), [
    ("That gap is noise. A tool that returns the higher one returns a confident, "
     "correctly-formatted answer about the ", FG),
    ("wrong store", RED, True), (" \u2014 and nobody ever finds out.", FG),
], accent=RED, size=15)

tf = textbox(s, Inches(0.7), Inches(4.0), Inches(11.95), Inches(1.8),
             align=PP_ALIGN.CENTER)
para(tf, "The dangerous failure isn't the one that crashes.",
     size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER, space_after=8, first=True)
para(tf, "It's the one that looks exactly like success.",
     size=28, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

# 7 · WHAT BIT US
s = slide(prs, """
100 seconds. If tight, do only #1 - it's the best story.
#1 lesson for the room: tool output is an API contract with a NON-DETERMINISTIC client.

Spare card if the room is infra-heavy: gemini-3.6-flash publishes only to `global`.
Our .env said us-central1, so every model call would have 404'd. Invisible until the
tests started loading .env.
""")
kicker(s, "Three things that actually bit us")
cw, gap = Inches(3.85), Inches(0.28)
card(s, Inches(0.7), Inches(1.15), cw, Inches(3.4), RED,
     "1 \u00b7 The model was eating base64", [
         ('"QLStkeuFHrk="  \u2192  5293.57', GREEN, True, MONO),
     ])
tf = s.shapes[-1].text_frame
rich(tf, [("Bigtable returned packed IEEE-754 doubles. Gemini didn't error \u2014 it "
           "quietly made things up. ", MUTED),
          ("~150 lines of \"working\" code never ran.", WHITE, True)], size=11)

card(s, Inches(0.7) + cw + gap, Inches(1.15), cw, Inches(3.4), AMBER,
     "2 \u00b7 Parallel wasn't parallel", [
         ("Sync tools run on the event loop unless a thread pool is configured \u2014 "
          "and adk web doesn't configure one. Two \"concurrent\" tools ran back to back.",
          MUTED),
     ])
tf = s.shapes[-1].text_frame
rich(tf, [("Nothing failed. It was just silently twice as slow.", WHITE, True)], size=11)

card(s, Inches(0.7) + 2 * (cw + gap), Inches(1.15), cw, Inches(3.4), PURPLE,
     "3 \u00b7 Iceberg rejects row-level security", [
         ("We chose an open table format for interoperability. It cost us "
          "BigQuery-native governance.", MUTED),
     ])
tf = s.shapes[-1].text_frame
rich(tf, [("A trade-off to make deliberately \u2014 not discover in production.",
           WHITE, True)], size=11)

tf = textbox(s, Inches(0.7), Inches(4.95), Inches(11.95), Inches(0.6),
             align=PP_ALIGN.CENTER)
rich(tf, [("Every one of these was ", MUTED), ("silent", AMBER, True),
          (". None threw an exception.", MUTED)],
     size=16, align=PP_ALIGN.CENTER, first=True)

# 8 · THE FLAW (setup)
s = slide(prs, """
DO NOT RUN LIVE - 42 seconds. Have it open in a second browser tab from before the talk.
"Here's one I ran earlier."

Let this slide sit for a beat. Let them accept it. THEN advance.
""")
kicker(s, "Where it's still wrong   ·   pre-baked   ·   42s")
title(s, "Live vs. 7-day cashier comparison", y=0.9, size=30)
table(s, Inches(0.7), Inches(2.1), Inches(11.95), [
    ["Metric",
     ("Live 1-hour", MUTED, SANS, PP_ALIGN.RIGHT),
     ("7-day baseline", MUTED, SANS, PP_ALIGN.RIGHT),
     ("Delta", MUTED, SANS, PP_ALIGN.RIGHT)],
    ["Manual Override Rate",
     ("71.05%", FG, MONO, PP_ALIGN.RIGHT),
     ("94.6%", FG, MONO, PP_ALIGN.RIGHT),
     ("\u221223.55pp", GREEN, MONO, PP_ALIGN.RIGHT)],
], col_widths=[Inches(4.45), Inches(2.5), Inches(2.5), Inches(2.5)], font_size=16)
tf = textbox(s, Inches(0.7), Inches(3.5), Inches(11.95), Inches(0.6))
para(tf, "Clean. Executive-ready. Two decimal places.",
     size=17, color=MUTED, first=True)

# 9 · THE REVEAL
s = slide(prs, """
This is your strongest slide. Showing a flaw in your own demo is the highest-trust move
available - a room of engineers assumes you're hiding something, and pre-empting it makes
everything else you claimed credible.

Verified against the raw table: 351 cashier_promo_abuse + 20 order_anomaly = 371.
""")
kicker(s, "Where it's still wrong")
tf = textbox(s, Inches(0.7), Inches(1.0), Inches(11.95), Inches(1.5),
             align=PP_ALIGN.CENTER)
para(tf, "94.6%", size=66, color=BLUE, bold=True, font=MONO,
     align=PP_ALIGN.CENTER, space_after=4, first=True)
rich(tf, [("is ", FG), ("not", RED, True), (" an override rate.", FG)],
     size=24, align=PP_ALIGN.CENTER)

cw = Inches(5.85)
card(s, Inches(0.7), Inches(3.35), cw, Inches(1.5), RED, "What it actually is", [
    ("351 promo-abuse alerts \u00f7 371 total alerts \u2014 the mix of this cashier's "
     "alert types.", MUTED),
])
card(s, Inches(0.7) + cw + Inches(0.25), Inches(3.35), cw, Inches(1.5), AMBER,
     "Why it can't be a rate", [
         ("Denominator is alerts, not transactions. The alerts table has no transaction "
          "count at all.", MUTED),
     ])
callout(s, Inches(0.7), Inches(5.15), Inches(11.95), Inches(0.95), [
    ("It compared a rate against a ratio and reported the difference to two decimals. "
     "Every number is real. ", FG),
    ("The subtraction is meaningless.", RED, True),
], accent=RED, size=15)

# 10 · SEMANTIC LAYER
s = slide(prs, """
Slow down here. This is the single most transferable idea in the talk, and it applies to
every NL2SQL / "chat with your data" product any of them will be asked to evaluate.
""")
kicker(s, "The lesson")
tf = textbox(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(2.5),
             align=PP_ALIGN.CENTER)
para(tf, "NL2SQL doesn't fail by writing broken SQL.",
     size=32, color=MUTED, align=PP_ALIGN.CENTER, space_after=14, first=True)
rich(tf, [("It fails by writing ", WHITE, True), ("valid", BLUE, True),
          (" SQL", WHITE, True)], size=36, align=PP_ALIGN.CENTER, space_after=8)
rich(tf, [("against a metric ", WHITE, True), ("nobody defined", BLUE, True),
          (".", WHITE, True)], size=36, align=PP_ALIGN.CENTER)
callout(s, Inches(1.2), Inches(4.7), Inches(10.95), Inches(1.5), [
    ("The semantic layer isn't documentation \u2014 it's the ", FG),
    ("control plane", GREEN, True),
    (".\nIf \"override rate\" isn't defined once, centrally, the model will define it "
     "for you. Differently. Every time.", FG),
], accent=GREEN, size=16)

# 11 · AUTHZ
s = slide(prs, """
If someone asks "what stops it querying data I'm not allowed to see" - this is the slide.

Note both layers are required: delegation without RLS just changes whose name is on the
query; RLS without delegation isolates nothing, because there's only ever one principal.
""")
kicker(s, "Where authorisation lives")
cw, gap = Inches(2.55), Inches(0.42)
for i, (top, bottom, accent) in enumerate([
    ("End user", "OAuth token", LINE),
    ("session.state", "app/auth.py", LINE),
    ("BigQuery job", "SESSION_USER()", LINE),
    ("Row access policy", "store_id IN grants", GREEN),
]):
    x = Inches(0.75) + i * (cw + gap)
    chip(s, x, Inches(1.15), cw, Inches(0.8), top, bottom, accent=accent)
    if i < 3:
        arrow(s, x + cw + Inches(0.02), Inches(1.32))

cw = Inches(5.85)
card(s, Inches(0.7), Inches(2.35), cw, Inches(1.7), RED, "A prompt is a request", [
    ('"Only show this manager their own store" \u2014 evaluated by a probabilistic '
     "system. Injection routes around it.", MUTED),
])
card(s, Inches(0.7) + cw + Inches(0.25), Inches(2.35), cw, Inches(1.7), GREEN,
     "A row access policy is a control", [
         ("Enforced by BigQuery beneath every tool. The model can be fully compromised "
          "and still not read another tenant's rows.", MUTED),
     ])
tf = textbox(s, Inches(0.7), Inches(4.5), Inches(11.95), Inches(1),
             align=PP_ALIGN.CENTER)
rich(tf, [("Proven, not assumed: granting one store collapsed visibility from ", FG),
          ("65,368", RED, True, MONO), (" rows to ", FG),
          ("1,233", GREEN, True, MONO), (".", FG)],
     size=17, align=PP_ALIGN.CENTER, first=True)

# 12 · TAKEAWAY
s = slide(prs, """
Three lines, then STOP. Don't trail off into "so yeah, that's it". End on line 3 and
take questions.
""")
kicker(s, "Takeaway")
for i, (accent, head, body) in enumerate([
    (BLUE, "1 \u00b7 The model is a router",
     "It holds no data and no permissions of its own. Everything hard about this was "
     "infrastructure."),
    (GREEN, "2 \u00b7 Authorisation belongs below the model",
     "Prompts are requests. Row access policies are controls."),
    (PURPLE, "3 \u00b7 Grounding and refusal are features",
     "The system that says \"I don't know which store you mean\" is the one you can put "
     "in front of an auditor."),
]):
    card(s, Inches(0.7), Inches(1.15) + i * Inches(1.35), Inches(11.95), Inches(1.15),
         accent, head, [(body, MUTED)])
tf = textbox(s, Inches(0.7), Inches(5.5), Inches(11.95), Inches(0.6),
             align=PP_ALIGN.CENTER)
para(tf, "37 tests  \u00b7  30 unit / 7 live integration  \u00b7  all green",
     size=13, color=MUTED, font=MONO, align=PP_ALIGN.CENTER, first=True)

# 13 · BACKUP: FALLBACK OUTPUT
s = slide(prs, "Use if the live demo fails. This is genuine captured output.")
kicker(s, "Backup \u00b7 Demo 1 output")
title(s, "ERR-PAY-4001 \u2014 EMV PIN Pad Tokenization Timeout", y=0.9, size=26)
tf = textbox(s, Inches(0.7), Inches(1.9), Inches(11.95), Inches(3.5))
for i, (n, txt) in enumerate([
    ("1.", "Do NOT re-swipe or re-charge immediately \u2014 prevents duplicate authorizations"),
    ("2.", "Reboot the payment module \u2014 hold Yellow + # for 3 seconds"),
    ("3.", "Check cable connections \u2014 12V PoweredUSB seated at both ends"),
    ("4.", "Manager Menu \u2192 Journal Audit Slip:  AUTHORIZED_UNSETTLED \u2192 print "
           "receipt  ·  VOIDED_ERROR \u2192 re-scan and re-present card"),
]):
    rich(tf, [(n + "  ", BLUE, True), (txt, FG)], size=16, space_after=14, first=(i == 0))
tf = textbox(s, Inches(0.7), Inches(5.3), Inches(11.95), Inches(0.5))
para(tf, "Source: Toshiba TCx 810 POS Hardware, Diagnostics & Service Guide",
     size=13, color=MUTED, first=True)

# 14 · BACKUP: NUMBERS
s = slide(prs, """
Numbers slide for Q&A. The HP/Toshiba point is worth volunteering if anyone asks how you
know retrieval quality is good - it shows you measured rather than assumed.
""")
kicker(s, "Backup \u00b7 Measured on this system")
table(s, Inches(0.7), Inches(1.0), Inches(11.95), [
    ["Prompt", ("Latency", MUTED, SANS, PP_ALIGN.RIGHT), "Tools dispatched"],
    ["Hardware error \u2192 runbook", ("9.3s", GREEN, MONO, PP_ALIGN.RIGHT),
     ("rag", MUTED, MONO)],
    ["\u201cthe Ginza store\u201d \u2192 refuses", ("5.3s", GREEN, MONO, PP_ALIGN.RIGHT),
     ("resolve", MUTED, MONO)],
    ["Live cashier metrics", ("15.7s", AMBER, MONO, PP_ALIGN.RIGHT),
     ("bigtable", MUTED, MONO)],
    ["Live vs 7-day comparison", ("42.4s", RED, MONO, PP_ALIGN.RIGHT),
     ("bigtable + analytics (parallel)", MUTED, MONO)],
    ["\u201cour Paris flagship\u201d \u2192 STORE_007", ("35.2s", RED, MONO, PP_ALIGN.RIGHT),
     ("resolve \u2192 analytics", MUTED, MONO)],
], col_widths=[Inches(5.0), Inches(2.0), Inches(4.95)], font_size=13)

cw, gap = Inches(3.85), Inches(0.28)
for i, (head, body) in enumerate([
    ("Retrieval corpus", "129 chunks \u00b7 5 docs \u00b7 768-dim"),
    ("Ranking caveat", "Pure cosine ranks the wrong vendor first: HP 0.7028 vs "
                       "Toshiba 0.6920"),
    ("Cost guardrails", "Byte cap \u00b7 row cap \u00b7 BigQuery custom quota"),
]):
    card(s, Inches(0.7) + i * (cw + gap), Inches(4.5), cw, Inches(1.3),
         BLUE, head, [(body, MUTED)])

# 15 · BACKUP: Q&A
s = slide(prs, "Don't present this - jump here during Q&A if useful.")
kicker(s, "Backup \u00b7 Likely questions")
table(s, Inches(0.7), Inches(1.0), Inches(11.95), [
    ["Question", "Answer"],
    [("Why not just write the SQL?", WHITE),
     "You'd write it four times for four questions. The value is routing across three "
     "storage systems behind one question. We did hand-write the retrieval SQL \u2014 "
     "only the warehouse layer is generated."],
    [("71% override rate? Absurd.", WHITE),
     "Correct \u2014 synthetic lab data. The plumbing is real; the numbers are generated."],
    [("How do you know retrieval is right?", WHITE),
     "We measure it. Reported score is true cosine, never inflated. Pure cosine ranks the "
     "wrong vendor first, so exact error-code match is a separate admission signal kept "
     "out of the score."],
    [("Why MCP, not a direct client?", WHITE),
     "Separate deployable, own IAM, own release cycle. We changed its query config without "
     "redeploying the agent."],
    [("Production ready?", WHITE),
     "No. End-user OAuth is unit-tested but never exercised with a real token; the semantic "
     "layer gap you just saw; no eval harness beyond 37 tests."],
], col_widths=[Inches(3.6), Inches(8.35)], font_size=11)

prs.save(OUT)
print(f"Wrote {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
